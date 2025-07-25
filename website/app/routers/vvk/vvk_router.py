from fastapi import APIRouter, HTTPException, Depends, Request, Form, Query
from fastapi.responses import HTMLResponse
from sqlalchemy.orm import Session
from typing import Optional
from typing import List
from ...model.database.database import SessionLocal
from ...model.vvk.models import VVK 
from ...model.vvk.schemas import VVK as vvk_schemas
from datetime import datetime, date
from fastapi.templating import Jinja2Templates
from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.responses import FileResponse
import zipfile
import os

from openpyxl import Workbook
import docx
import tempfile
import shutil
import re
import os

import logging

logging.basicConfig(level=logging.INFO)

router = APIRouter()


frontend_dir = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "../templates"))

templates = Jinja2Templates(directory=frontend_dir)

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()
"""
@router.post("/vvk/", response_model=vvk_schemas.VVK)
def create_vvk(vvk_data: vvk_schemas.VVKCreate, db: Session = Depends(get_db)):
    db_vvk = VVK.VVK(**vvk_data.dict())
    db.add(db_vvk)
    db.commit()
    db.refresh(db_vvk)
    return db_vvk

@router.get("/vvk/{vvk_id}", response_model=vvk_schemas.VVK)
def read_vvk(vvk_id: int, db: Session = Depends(get_db)):
    db_vvk = db.query(VVK.VVK).filter(VVK.VVK.id == vvk_id).first()
    if db_vvk is None:
        raise HTTPException(status_code=404, detail="VVK not found")
    return db_vvk

@router.put("/vvk/{vvk_id}", response_model=vvk_schemas.VVK)
def update_vvk(vvk_id: int, vvk_data: vvk_schemas.VVKUpdate, db: Session = Depends(get_db)):
    db_vvk = db.query(VVK.VVK).filter(VVK.VVK.id == vvk_id).first()
    if db_vvk is None:
        raise HTTPException(status_code=404, detail="VVK not found")
    for key, value in vvk_data.dict(exclude_unset=True).items():
        setattr(db_vvk, key, value)
    db.commit()
    db.refresh(db_vvk)
    return db_vvk

@router.delete("/vvk/{vvk_id}", response_model=vvk_schemas.VVK)
def delete_vvk(vvk_id: int, db: Session = Depends(get_db)):
    db_vvk = db.query(VVK.VVK).filter(VVK.VVK.id == vvk_id).first()
    if db_vvk is None:
        raise HTTPException(status_code=404, detail="VVK not found")
    db.delete(db_vvk)
    db.commit()
    return db_vvk
"""

@router.get("/", response_class=HTMLResponse)
async def read_documents(request: Request):   
    return templates.TemplateResponse("vvk/vvk_table.html", {"request": request})

save_directory = "/temp"
os.makedirs(save_directory, exist_ok=True)  # Создайте директорию, если она не существует

@router.post("/upload/")
async def create_upload_files(files: List[UploadFile] = File(...)):
    try:
        wb = Workbook()
        ws = wb.active
        ws.append(["Воинская часть", "Воинское звание", "ФИО", "Дата рождения", "Дата оформления заключения", "Заключение", "Номер"])
        
        for file in files:
            # Сохраняем файлы DOCX на диск
            doc_path = os.path.join(save_directory, file.filename)
            with open(doc_path, "wb") as f:
                shutil.copyfileobj(file.file, f)
            logging.debug(f"DOCX file saved at: {doc_path}")
            
            # Читаем и обрабатываем DOCX файл
            doc = docx.Document(doc_path)
            data = extract_data_from_text(doc)
            ws.append([data[key] for key in data])

        output_path = os.path.join(save_directory, "result.xlsx")
        wb.save(output_path)
        logging.debug(f"Excel file saved at: {output_path}")

        if not os.path.exists(output_path):
            logging.error("Excel file does not exist.")
            raise HTTPException(status_code=500, detail="Failed to save the Excel file.")

        return FileResponse(output_path, filename="result.xlsx", media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:
        logging.error(f"An error occurred: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))



def save_to_excel(data, output_file):
    wb = Workbook()
    ws = wb.active

    for entry in data:
        ws.append([
            entry["Воинская часть"], entry["Воинское звание"], entry["ФИО"],
            entry["Дата рождения"], entry["Дата оформления заключения"],
            entry["Заключение"], entry["Номер"]
        ])

    wb.save(output_file)




def extract_data_from_text(doc):
    text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
    lines = text.split('\n')
    data = {
        "Воинская часть": None,
        "Воинское звание": None,
        "ФИО": None,
        "Дата рождения": None,
        "Дата оформления заключения": None,
        "Заключение": None,
        "Номер": None
    }
    collecting_conclusion = False
    conclusion_lines = []

    for line in lines:
        if "комиссии №" in line:
            # Extract the number assuming it follows "комиссии №"
            number_start = line.find("комиссии №") + len("комиссии №")
            data["Номер"] = line[number_start:].strip().split(' ')[0]  
        if "Фамилия," in line:
            data["ФИО"] = line.split(':')[-1].strip()
        if "Дата рождения:" in line:
            data["Дата рождения"] = line.split(':')[-1].strip()
        if "Воинское звание:" in line:
            data["Воинское звание"] = line.split(':')[-1].strip()
        if "Место службы: в/ч" in line:
            data["Воинская часть"] = line.split('в/ч')[-1].strip()
        if "военно-врачебной комиссии" in line and "«" in line:
            match = re.search(r'«(\d{1,2})»\s*([а-яА-Я]+)\s*(\d{4})\s*г\.', line)
            if match:
                data["Дата оформления заключения"] = f'«{match.group(1)}» {match.group(2)} {match.group(3)} г.'
        if "Заключение военно-врачебной комиссии:" in line or "Заключение военно-врачебной комиссии" in line:
            collecting_conclusion = True
            continue
        if "ВрИО " in line or "Председатель" in line:
            break
        if collecting_conclusion:
            conclusion_lines.append(line.strip())

    data["Заключение"] = ' '.join(conclusion_lines).strip() 
    return data